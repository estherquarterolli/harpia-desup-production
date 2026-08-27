# -*- coding: utf-8 -*-
"""
extract_disciplinas.py
Extracts curricular component data (codigo, nome, sigla, carga_horaria) from
PDF and DOCX files and generates:
  - disciplinas_encontradas.txt  (plain-text summary)
  - update_component_codes.sql   (SQL UPDATE statements)
"""

import re
import os
import sys
from pathlib import Path

BASE_DIR = Path(r"C:\Users\esther.santos\AppData\Local\Temp\Rar$DRa9068.19783")
OUT_DIR  = Path(r"C:\Users\esther.santos\Documents\GitHub\AllocGest-DESUP")

# ── file list in priority order ───────────────────────────────────────────────
FILES = [
    # DOCX first (priority)
    "Matrizes-Desup (COMPILADO-DISCIPLINAS).docx",
    "PPPC_FAETERJ_Petropolis_TIC_v7.docx",
    "Projeto Pedagógico de Curso - PED - FAETERJ Bom Jesus.docx",
    "Projeto Pedagógico de Curso - PED - FAETERJ Itaperuna.docx",
    "Projeto Pedagógico de Curso - PED - FAETERJ Pádua (1).docx",
    "Projeto Pedagógico de Curso - PED - FAETERJ Pádua.docx",
    "Projeto Pedagógico de Curso - PED - FAETERJ Três Rios.docx",
    "Projeto Pedagógico de Curso - PED - ISEPAM.docx",
    # PDFs
    "Matrizes-Desup (COMPILADO-DISCIPLINAS).pdf",
    "03_01_2024___Projeto_Pedagogico_de_Curso___ADS_Paracambi__1_.pdf",
    "Modelo___Projeto_Pedagogico_de_Curso___TGA.pdf",
    "PPC do Curso de Pedagogia (Licenciatura) - abril de 2023.pdf",
    "PPC_ADS_FAETERJ_Rio_v2_25Jan2024.pdf",
    "Projeto_Pedagogico_de_Curso___Gestao_Portuaria.pdf",
    "Projeto_Pedagogico_de_Curso___TPG___versao_final_15_12_23.pdf",
    "Projeto_Pedagogico_de_Curso_ADS_Barra_Mansa_2024_rev_v4_ERRATA.pdf",
]

# ── regex patterns ────────────────────────────────────────────────────────────
# Código like MA-101, PO-201, ADS-001, INF-302 …
CODE_RE = re.compile(
    r'\b([A-Z]{2,6}[-–]?\d{2,4})\b'
)
# Carga horária: digits followed by h / hs / horas / hora-relógio
CH_RE = re.compile(
    r'\b(\d{1,4})\s*(?:h\.?r?\.?|hs?\.?|horas?|hora[s\-]rel[oó]gio)\b',
    re.IGNORECASE,
)
# Column header synonyms
HEADER_CODIGO   = re.compile(r'c[oó]d(igo)?\.?|code', re.IGNORECASE)
HEADER_NOME     = re.compile(r'nome|disciplina|componente|unidade curricular', re.IGNORECASE)
HEADER_SIGLA    = re.compile(r'sigla|abrev', re.IGNORECASE)
HEADER_CH       = re.compile(r'ch|carga.?hor[aá]ria|horas?|total.?h', re.IGNORECASE)


# ── helpers ───────────────────────────────────────────────────────────────────

def clean(s):
    if s is None:
        return ""
    return re.sub(r'\s+', ' ', str(s)).strip()


def looks_like_code(s):
    s = clean(s)
    # 2-8 chars, has letters and digits OR letter-dash-digits
    if not s:
        return False
    if CODE_RE.fullmatch(s):
        return True
    if re.match(r'^[A-Z]{2,6}\d{1,4}$', s):
        return True
    return False


def extract_ch(s):
    """Return first number found in a string that looks like a workload value."""
    s = clean(s)
    m = CH_RE.search(s)
    if m:
        return m.group(1)
    # bare number between 20 and 600
    m2 = re.search(r'\b(\d{2,3})\b', s)
    if m2:
        v = int(m2.group(1))
        if 20 <= v <= 600:
            return str(v)
    return ""


def map_columns(header_row):
    """Return dict: {'codigo': idx, 'nome': idx, 'sigla': idx, 'ch': idx}"""
    mapping = {}
    for i, cell in enumerate(header_row):
        c = clean(cell)
        if HEADER_CODIGO.search(c) and 'codigo' not in mapping:
            mapping['codigo'] = i
        elif HEADER_NOME.search(c) and 'nome' not in mapping:
            mapping['nome'] = i
        elif HEADER_SIGLA.search(c) and 'sigla' not in mapping:
            mapping['sigla'] = i
        elif HEADER_CH.search(c) and 'ch' not in mapping:
            mapping['ch'] = i
    return mapping


def row_to_record(row, col_map):
    """Convert a table row to a discipline dict using column mapping."""
    def get(key):
        idx = col_map.get(key)
        if idx is not None and idx < len(row):
            return clean(row[idx])
        return ""

    record = {
        'codigo': get('codigo'),
        'nome':   get('nome'),
        'sigla':  get('sigla'),
        'ch':     get('ch'),
    }
    # if no explicit code col, scan all cells for a code pattern
    if not record['codigo']:
        for cell in row:
            c = clean(cell)
            if looks_like_code(c):
                record['codigo'] = c
                break
    # normalise CH
    if record['ch']:
        record['ch'] = extract_ch(record['ch']) or record['ch']

    return record


def is_valid_record(r):
    nome = r.get('nome', '').strip()
    return (
        len(nome) >= 4
        and not HEADER_NOME.fullmatch(nome)   # skip header rows stored as data
        and nome.lower() not in ('disciplina', 'nome', 'componente curricular', '')
    )


# ── DOCX extraction ───────────────────────────────────────────────────────────

def extract_docx(filepath):
    try:
        from docx import Document
    except ImportError:
        print("  [WARN] python-docx not installed; skipping DOCX.")
        return []

    records = []
    doc = Document(filepath)

    for table in doc.tables:
        if not table.rows:
            continue

        # Try first row as header
        header_row = [clean(c.text) for c in table.rows[0].cells]
        col_map = map_columns(header_row)

        # Need at least a nome column to proceed with mapping
        if 'nome' in col_map:
            for row in table.rows[1:]:
                cells = [clean(c.text) for c in row.cells]
                r = row_to_record(cells, col_map)
                if is_valid_record(r):
                    records.append(r)
        else:
            # Heuristic: any row where a cell is a code and another is long text
            for row in table.rows:
                cells = [clean(c.text) for c in row.cells]
                code_cells  = [c for c in cells if looks_like_code(c)]
                long_cells  = [c for c in cells if len(c) > 10 and not looks_like_code(c)]
                if code_cells and long_cells:
                    ch_cells = [c for c in cells if extract_ch(c)]
                    r = {
                        'codigo': code_cells[0],
                        'nome':   long_cells[0],
                        'sigla':  '',
                        'ch':     extract_ch(ch_cells[0]) if ch_cells else '',
                    }
                    if is_valid_record(r):
                        records.append(r)

    # Also scan paragraphs for inline discipline lines like:
    # "MA-101 – Matemática Básica   60h"
    para_pattern = re.compile(
        r'([A-Z]{2,6}[-–]\d{2,4})\s*[-–:]\s*(.{5,80?}?)\s+(\d{2,3})\s*h',
        re.IGNORECASE,
    )
    for para in doc.paragraphs:
        text = clean(para.text)
        for m in para_pattern.finditer(text):
            r = {
                'codigo': m.group(1).upper(),
                'nome':   clean(m.group(2)),
                'sigla':  '',
                'ch':     m.group(3),
            }
            if is_valid_record(r):
                records.append(r)

    return records


# ── PDF extraction ────────────────────────────────────────────────────────────

def extract_pdf(filepath):
    try:
        import pdfplumber
    except ImportError:
        print("  [WARN] pdfplumber not installed; skipping PDF.")
        return []

    records = []

    try:
        with pdfplumber.open(filepath) as pdf:
            for page in pdf.pages:
                # ── tables ──
                tables = page.extract_tables() or []
                for table in tables:
                    if not table:
                        continue
                    header = [clean(c) for c in table[0]]
                    col_map = map_columns(header)

                    if 'nome' in col_map:
                        for row in table[1:]:
                            r = row_to_record([clean(c) for c in row], col_map)
                            if is_valid_record(r):
                                records.append(r)
                    else:
                        for row in table:
                            cells = [clean(c) for c in row]
                            code_cells = [c for c in cells if looks_like_code(c)]
                            long_cells = [c for c in cells if len(c) > 10 and not looks_like_code(c)]
                            if code_cells and long_cells:
                                ch_cells = [c for c in cells if extract_ch(c)]
                                r = {
                                    'codigo': code_cells[0],
                                    'nome':   long_cells[0],
                                    'sigla':  '',
                                    'ch':     extract_ch(ch_cells[0]) if ch_cells else '',
                                }
                                if is_valid_record(r):
                                    records.append(r)

                # ── plain text scan ──
                text = page.extract_text() or ""
                para_pattern = re.compile(
                    r'([A-Z]{2,6}[-–]\d{2,4})\s*[-–:]\s*(.{5,80}?)\s+(\d{2,3})\s*h',
                )
                for m in para_pattern.finditer(text):
                    r = {
                        'codigo': m.group(1).upper(),
                        'nome':   clean(m.group(2)),
                        'sigla':  '',
                        'ch':     m.group(3),
                    }
                    if is_valid_record(r):
                        records.append(r)
    except Exception as e:
        print(f"  [ERROR] {filepath.name}: {e}")

    return records


# ── deduplication ─────────────────────────────────────────────────────────────

def dedup(records):
    seen = {}
    out = []
    for r in records:
        key = r['nome'].lower().strip()
        if key not in seen:
            seen[key] = True
            out.append(r)
    return out


# ── SQL generation ────────────────────────────────────────────────────────────

def escape_sql(s):
    return s.replace("'", "''")


def make_sql(filename, course, record):
    nome   = escape_sql(record['nome'])
    codigo = escape_sql(record.get('codigo', ''))
    sigla  = escape_sql(record.get('sigla', ''))
    ch     = escape_sql(record.get('ch', ''))

    sets = []
    if codigo:
        sets.append(f"codigo = '{codigo}'")
    if sigla:
        sets.append(f"sigla = '{sigla}'")
    if ch:
        sets.append(f"carga_horaria = {ch}" if ch.isdigit() else f"carga_horaria = '{ch}'")

    if not sets:
        sets.append("-- (no code/sigla/ch found)")

    set_clause = ", ".join(sets)
    return (
        f"-- From: {filename} | Course: {course}\n"
        f"UPDATE courses_curricularcomponent SET {set_clause} "
        f"WHERE nome ILIKE '{nome}';\n"
    )


# ── course name heuristic ─────────────────────────────────────────────────────

def guess_course(filename):
    fn = filename.lower()
    if 'ads' in fn or 'analise' in fn or 'sistemas' in fn:
        return 'ADS - Análise e Desenvolvimento de Sistemas'
    if 'pedagogia' in fn or 'ped' in fn:
        return 'Pedagogia'
    if 'gestao_portuaria' in fn or 'portuaria' in fn:
        return 'Gestão Portuária'
    if 'tga' in fn:
        return 'TGA - Tecnologia em Gestão Ambiental'
    if 'tpg' in fn:
        return 'TPG - Tecnologia em Processos Gerenciais'
    if 'tic' in fn:
        return 'TIC - Tecnologia da Informação e Comunicação'
    if 'compilado' in fn or 'matrizes' in fn:
        return '(Compilado Geral)'
    return '(Desconhecido)'


# ── main ──────────────────────────────────────────────────────────────────────

def main():
    txt_path = OUT_DIR / "disciplinas_encontradas.txt"
    sql_path = OUT_DIR / "update_component_codes.sql"

    all_sql_lines = [
        "-- Auto-generated by extract_disciplinas.py\n",
        "-- Review before running against the database!\n\n",
    ]
    txt_lines = ["DISCIPLINAS ENCONTRADAS\n", "=" * 70 + "\n\n"]

    total_found = 0
    files_no_data = []
    code_patterns_seen = set()

    for filename in FILES:
        filepath = BASE_DIR / filename
        if not filepath.exists():
            print(f"[MISSING] {filename}")
            txt_lines.append(f"[ARQUIVO NÃO ENCONTRADO] {filename}\n\n")
            continue

        ext = filepath.suffix.lower()
        print(f"\nProcessing: {filename}")

        if ext == '.docx':
            records = extract_docx(filepath)
        elif ext == '.pdf':
            records = extract_pdf(filepath)
        else:
            print(f"  [SKIP] Unknown extension: {ext}")
            continue

        records = dedup(records)
        course  = guess_course(filename)
        count   = len(records)
        total_found += count

        print(f"  -> {count} discipline(s) found")

        if count == 0:
            files_no_data.append(filename)
            txt_lines.append(f"FILE: {filename}\n")
            txt_lines.append(f"  Course: {course}\n")
            txt_lines.append(f"  [SEM DADOS – pode ser PDF escaneado ou formato não reconhecido]\n\n")
            continue

        txt_lines.append(f"FILE: {filename}\n")
        txt_lines.append(f"  Course: {course}\n")
        txt_lines.append(f"  Disciplines found: {count}\n")

        for r in records:
            # collect code patterns
            if r.get('codigo'):
                m = CODE_RE.match(r['codigo'])
                if m:
                    prefix = re.sub(r'\d', 'N', r['codigo'])
                    code_patterns_seen.add(prefix)

            line = f"    • [{r.get('codigo','?')}] {r['nome']}"
            if r.get('sigla'):
                line += f" ({r['sigla']})"
            if r.get('ch'):
                line += f"  CH={r['ch']}h"
            txt_lines.append(line + "\n")

            sql = make_sql(filename, course, r)
            all_sql_lines.append(sql)

        txt_lines.append("\n")

    # ── summary ──────────────────────────────────────────────────────────────
    summary = (
        "\n" + "=" * 70 + "\n"
        f"TOTAL: {total_found} disciplines across {len(FILES)} files\n"
        f"Code patterns detected: {', '.join(sorted(code_patterns_seen)) or 'none'}\n"
        f"Files with no extractable data ({len(files_no_data)}):\n"
    )
    for f in files_no_data:
        summary += f"  - {f}\n"

    txt_lines.append(summary)
    print(summary)

    # ── write output files ────────────────────────────────────────────────────
    with open(txt_path, "w", encoding="utf-8") as f:
        f.writelines(txt_lines)
    print(f"\nWrote: {txt_path}")

    with open(sql_path, "w", encoding="utf-8") as f:
        f.writelines(all_sql_lines)
    print(f"Wrote: {sql_path}")


if __name__ == "__main__":
    main()
